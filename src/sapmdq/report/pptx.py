"""Vorlage für die Ergebnispräsentation (FA-705).

Erzeugt eine PowerPoint-Datei mit vorbelegten Kernkennzahlen als
Ausgangspunkt für die Ergebnisbesprechung. Die Anforderung ist mit "Could"
eingestuft; entsprechend ist ``python-pptx`` eine optionale Abhängigkeit.
Fehlt sie, entfällt der Export mit einem Hinweis, statt den Lauf zu
gefährden.

Bewusst enthält die Vorlage keine Befunddetails, sondern nur verdichtete
Zahlen und den Vorbehalt zur Aussagekraft. Eine Präsentation wird
weitergereicht, oft über den Kreis hinaus, der die Daten sehen darf (DS-02).
"""

from __future__ import annotations

from pathlib import Path

from sapmdq.logging_setup import get_logger
from sapmdq.results import RunResult

logger = get_logger("report.pptx")


def is_available() -> bool:
    """Prüft, ob die optionale Abhängigkeit vorhanden ist."""
    try:
        import pptx  # noqa: F401
    except ImportError:
        return False
    return True


def write_presentation(result: RunResult, target: Path) -> Path | None:
    """Erzeugt die Präsentationsvorlage."""
    if not is_available():
        logger.info(
            "python-pptx ist nicht installiert - die Präsentationsvorlage entfällt. "
            "Nachinstallation: pip install 'sapmdq[pptx]'"
        )
        return None

    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    title_layout = presentation.slide_layouts[0]
    bullet_layout = presentation.slide_layouts[1]

    # ------------------------------------------------------- Titelfolie
    slide = presentation.slides.add_slide(title_layout)
    slide.shapes.title.text = "Prüfung der SAP-Stammdaten"
    slide.placeholders[1].text = (
        f"{result.config.project.name}\n"
        f"{result.config.project.customer}\n"
        f"Lauf {result.run_id} | Quellsystem {result.config.project.source_system}"
    )

    # ---------------------------------------------- Aussagekraft zuerst
    slide = presentation.slides.add_slide(bullet_layout)
    slide.shapes.title.text = "Aussagekraft dieses Ergebnisses"
    frame = slide.placeholders[1].text_frame
    frame.text = result.coverage.qualification() if result.coverage else "Keine Angabe."
    frame.paragraphs[0].font.size = Pt(16)
    frame.word_wrap = True

    # ---------------------------------------------------- Kennzahlen
    slide = presentation.slides.add_slide(bullet_layout)
    slide.shapes.title.text = "Kennzahlen"
    frame = slide.placeholders[1].text_frame
    executed = len([e for e in result.all_executions if e.status.value == "ausgefuehrt"])
    entries = [
        f"Verarbeitete Sätze: {result.rows_ingested:,}".replace(",", "."),
        f"Ausgeführte Prüfungen: {executed} von {result.coverage.total if result.coverage else 0}",
        f"Coverage-Grad: {result.coverage.coverage_ratio:.0%}" if result.coverage else "",
        f"Befunde gesamt: {result.effective_findings}",
    ]
    if result.score and result.score.overall is not None:
        entries.append(f"Data-Quality-Score: {result.score.overall} von 100")
    if result.delta:
        entries.append(result.delta.summary_line())

    frame.text = entries[0]
    for entry in entries[1:]:
        if not entry:
            continue
        paragraph = frame.add_paragraph()
        paragraph.text = entry
        paragraph.level = 0

    # ----------------------------------------------- Befunde je Kategorie
    if result.score:
        slide = presentation.slides.add_slide(bullet_layout)
        slide.shapes.title.text = "Datenqualität je Objektbereich"
        frame = slide.placeholders[1].text_frame
        first = True
        for area in result.score.areas:
            text = (
                f"{area.area}: {area.score if area.score is not None else 'nicht bewertbar'} "
                f"({area.grade}, {area.findings} Befunde)"
            )
            if first:
                frame.text = text
                first = False
            else:
                frame.add_paragraph().text = text

    # ------------------------------------------------ Nachforderung
    if result.coverage and result.coverage.demand_list:
        slide = presentation.slides.add_slide(bullet_layout)
        slide.shapes.title.text = "Was zusätzliche Lieferungen bringen"
        frame = slide.placeholders[1].text_frame
        first = True
        for candidate in result.coverage.demand_list[:6]:
            text = f"{candidate.request}: +{candidate.direct_count} Prüfungen"
            if first:
                frame.text = text
                first = False
            else:
                frame.add_paragraph().text = text

    # ------------------------------------------------ Hinweis zur Vorlage
    slide = presentation.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(8.4), Inches(2.5))
    box.text_frame.word_wrap = True
    box.text_frame.text = (
        "Diese Datei ist eine Vorlage mit vorbelegten Kennzahlen und keine fertige "
        "Präsentation. Befunddetails sind bewusst nicht enthalten: eine Präsentation "
        "wird weitergereicht, oft über den Kreis hinaus, der die Daten sehen darf."
    )

    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))
    logger.info("Präsentationsvorlage geschrieben: %s", target)
    return target
